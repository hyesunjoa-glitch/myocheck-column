"""GEO 측정 결과를 시트 3장(대시보드 / 원본데이터 / 질문세트 v1.0) 엑셀로 만든다.

사용: python3 -I scripts/geo_build_sheet.py <runs.json> <question-set.csv> <out.xlsx>
기준서: docs/geo-standard.md 4장 ①
"""
import csv
import json
import re
import sys
from collections import Counter
from datetime import datetime, timedelta, timezone
from urllib.parse import urlparse

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

RUNS, QCSV, OUT = sys.argv[1], sys.argv[2], sys.argv[3]
R = json.load(open(RUNS, encoding="utf-8"))
QS = list(csv.DictReader(open(QCSV, encoding="utf-8")))
KST = timezone(timedelta(hours=9))

STRICT = re.compile(r"묘책\s?사주|myocheck|MYO\s*:?\s*CHECK|재회막차", re.I)
COMPS = {
    "청월당": ["청월당"], "타이트사주": ["타이트사주", "타이트 사주"], "포스텔러": ["포스텔러"],
    "점신": ["점신"], "헬로우봇": ["헬로우봇"], "사주GPT": ["사주GPT", "사주 GPT"],
    "묘한사주": ["묘한사주"], "식스텔링": ["식스텔링"], "사주아이": ["사주아이"],
    "용용사주": ["용용사주"], "홍카페": ["홍카페"], "천을귀인": ["천을귀인"],
    "운세닷컴": ["운세닷컴"], "사주플러스": ["사주플러스"], "평생사주": ["평생사주"],
    "사주나루": ["사주나루"], "홍연": ["홍연"], "사주풀이도우미": ["사주풀이도우미"],
    # 측정 전 목록에 없었지만 답을 읽어 보니 반복해서 나온 온라인 서비스 (2026-10-19 추가)
    "이웃집 마녀": ["이웃집 마녀", "이웃집마녀"], "사주AI 소연": ["사주AI 소연", "사주 AI 소연"],
}
IMG = "images.openai.com"


def dom(u):
    d = urlparse(u).netloc.lower()
    return d[4:] if d.startswith("www.") else d


def ours(d):
    return d == "myocheck.kr" or d.endswith(".myocheck.kr")


def kst(iso):
    return datetime.fromisoformat(iso.replace("Z", "+00:00")).astimezone(KST)


# ---------- 공통 스타일 ----------
F = "Arial"
f_base = Font(name=F, size=10)
f_bold = Font(name=F, size=10, bold=True)
f_head = Font(name=F, size=10, bold=True, color="FFFFFF")
f_title = Font(name=F, size=16, bold=True)
f_h2 = Font(name=F, size=12, bold=True)
f_big = Font(name=F, size=20, bold=True)
f_note = Font(name=F, size=9, color="555555")
fill_head = PatternFill("solid", fgColor="3B3B58")
fill_card = PatternFill("solid", fgColor="F3F1FA")
fill_warn = PatternFill("solid", fgColor="FFF4E5")
thin = Side(style="thin", color="D0D0D8")
box = Border(left=thin, right=thin, top=thin, bottom=thin)
wrap = Alignment(wrap_text=True, vertical="top")
center = Alignment(horizontal="center", vertical="center", wrap_text=True)

wb = Workbook()
ws_d = wb.active
ws_d.title = "대시보드"
ws_r = wb.create_sheet("원본데이터")
ws_q = wb.create_sheet("질문세트 v1.0")

# ---------- 원본데이터 ----------
qmap = {q["id"]: q for q in QS}
order = {q["id"]: i for i, q in enumerate(QS)}
rows = sorted(R.values(), key=lambda v: (order[v["ID"]], v["RUN"]))
rhead = ["실행ID", "질문ID", "회차", "단계", "질문에 묘책 이름", "질문", "측정 시각(한국)", "모델",
         "웹검색 함", "답에 묘책 언급", "우리 사이트 인용", "우리 사이트 검색결과 노출",
         "인용 출처 도메인(이미지 제외)", "경쟁사 언급", "답 길이(자)", "답 원문", "대화ID",
         "계산용_인용도메인", "계산용_경쟁사"]
ws_r.append(rhead)
for v in rows:
    q = qmap[v["ID"]]
    cd = sorted({dom(u) for u in v["cited"]} - {IMG})
    sd = {dom(u) for u in v["seen"]}
    comp = [c for c, al in COMPS.items() if any(a in v["text"] for a in al)]
    ws_r.append([
        f'{v["ID"]}-{v["RUN"]}', v["ID"], v["RUN"], q["단계"], q["브랜드"], v["Q"],
        kst(v["ts"]).strftime("%Y-%m-%d %H:%M"), v["model"],
        "Y" if v["seen"] else "N",
        "Y" if STRICT.search(v["text"]) else "N",
        "Y" if any(ours(d) for d in cd) else "N",
        "Y" if any(ours(d) for d in sd) else "N",
        ", ".join(cd), ", ".join(comp), len(v["text"]), v["text"], v["cid"],
        "|" + "|".join(cd) + "|", "|" + "|".join(comp) + "|",
    ])
NR = len(rows) + 1
widths = [9, 7, 5, 9, 9, 40, 16, 11, 8, 8, 9, 10, 40, 22, 9, 60, 38, 20, 20]
for i, w in enumerate(widths, 1):
    ws_r.column_dimensions[get_column_letter(i)].width = w
for row in ws_r.iter_rows(min_row=2, max_row=NR):
    for c in row:
        c.font = f_base
        c.alignment = Alignment(vertical="top", wrap_text=c.column in (6, 13, 14))
    row[15].alignment = Alignment(vertical="top", wrap_text=False)
for c in ws_r[1]:
    c.font = f_head
    c.fill = fill_head
    c.alignment = center
ws_r.column_dimensions["R"].hidden = True
ws_r.column_dimensions["S"].hidden = True
ws_r.freeze_panes = "C2"
t = Table(displayName="raw", ref=f"A1:{get_column_letter(len(rhead))}{NR}")
t.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=True)
ws_r.add_table(t)

# 원본데이터 열 위치 (수식에서 씀)
RD = "원본데이터"
col = {h: get_column_letter(i + 1) for i, h in enumerate(rhead)}


def rng(h):
    return f"{RD}!${col[h]}$2:${col[h]}${NR}"


# ---------- 질문세트 v1.0 ----------
qhead = list(QS[0].keys()) + ["측정 회차", "묘책 언급 회차", "우리 사이트 인용 회차",
                               "웹검색 회차", "상태", "다음 달 글감 순위"]
ws_q.append(qhead)
for q in QS:
    ws_q.append([q[k] for k in QS[0].keys()])
NQ = len(QS) + 1
base = len(QS[0])
L = {h: get_column_letter(base + i + 1) for i, h in enumerate(qhead[base:])}
for r in range(2, NQ + 1):
    ws_q[f'{L["측정 회차"]}{r}'] = f"=COUNTIF({rng('질문ID')},$A{r})"
    ws_q[f'{L["묘책 언급 회차"]}{r}'] = f'=COUNTIFS({rng("질문ID")},$A{r},{rng("답에 묘책 언급")},"Y")'
    ws_q[f'{L["우리 사이트 인용 회차"]}{r}'] = f'=COUNTIFS({rng("질문ID")},$A{r},{rng("우리 사이트 인용")},"Y")'
    ws_q[f'{L["웹검색 회차"]}{r}'] = f'=COUNTIFS({rng("질문ID")},$A{r},{rng("웹검색 함")},"Y")'
    m, b, s = (f'{L["측정 회차"]}{r}', f'{L["묘책 언급 회차"]}{r}', f'{L["웹검색 회차"]}{r}')
    ws_q[f'{L["상태"]}{r}'] = (
        f'=IF({m}=0,"미측정",IF($D{r}="O",IF({L["우리 사이트 인용 회차"]}{r}>0,'
        f'"이름 질문 · 우리 사이트 인용됨","이름 질문 · 인용 안 됨"),IF({b}>0,"노출","빈칸")))')
    ws_q[f'{L["다음 달 글감 순위"]}{r}'] = (
        f'=IF(AND($D{r}="X",{m}>0,{b}=0),IF({s}>0,"1순위(ChatGPT가 검색함)","2순위(검색 안 함)"),"")')
qw = [6, 9, 44, 6, 9, 11, 22, 11, 44, 8, 9, 10, 8, 22, 20]
for i, w in enumerate(qw, 1):
    ws_q.column_dimensions[get_column_letter(i)].width = w
for row in ws_q.iter_rows(min_row=2, max_row=NQ):
    for c in row:
        c.font = f_base
        c.alignment = Alignment(vertical="top", wrap_text=c.column in (3, 9, 14, 15))
for c in ws_q[1]:
    c.font = f_head
    c.fill = fill_head
    c.alignment = center
ws_q.freeze_panes = "D2"
t2 = Table(displayName="qset", ref=f"A1:{get_column_letter(len(qhead))}{NQ}")
t2.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=True)
ws_q.add_table(t2)
QD = "'질문세트 v1.0'"


def qr(h):
    letter = L[h] if h in L else get_column_letter(qhead.index(h) + 1)
    return f"{QD}!${letter}$2:${letter}${NQ}"


# ---------- 대시보드 (세로 한 줄 흐름: A~D 4칸만 사용) ----------
d = ws_d
for c, w in zip("ABCD", [38, 13, 24, 58]):
    d.column_dimensions[c].width = w
ts = sorted(kst(v["ts"]) for v in R.values())
period = f'{ts[0]:%Y-%m-%d %H:%M} ~ {ts[-1]:%Y-%m-%d %H:%M} (한국 시간)'
row = 1


def h2(text):
    global row
    row += 1
    d[f"A{row}"] = text
    d[f"A{row}"].font = f_h2
    row += 1


def header(cols):
    global row
    for j, h in enumerate(cols):
        c = d.cell(row=row, column=1 + j, value=h)
        c.font = f_head
        c.fill = fill_head
        c.alignment = center
    row += 1


def line(vals, bold_cols=(), fmt=None, wrap_cols=(3,), height=None, fill=None):
    global row
    for j, v in enumerate(vals):
        c = d.cell(row=row, column=1 + j, value=v)
        c.font = f_bold if j in bold_cols else f_base
        c.border = box
        c.alignment = Alignment(vertical="top", wrap_text=j in wrap_cols)
        if fill:
            c.fill = fill
        if fmt and j in fmt:
            c.number_format = fmt[j]
    if height:
        d.row_dimensions[row].height = height
    row += 1


def para(text, font=f_base, height=30):
    global row
    d[f"A{row}"] = text
    d[f"A{row}"].font = font
    d[f"A{row}"].alignment = wrap
    d.merge_cells(f"A{row}:D{row}")
    d.row_dimensions[row].height = height
    row += 1


d["A1"] = "묘책 GEO 측정 대시보드 — 1차 (ChatGPT)"
d["A1"].font = f_title
d["A2"] = f"측정 기간 {period} · 질문세트 v1.0 · ChatGPT 임시 채팅 · 질문당 3회"
d["A2"].font = f_note
row = 3

h2("핵심 숫자 3개")
header(["지표", "값", "질문 / 회차", "뜻"])
n_meas = f'COUNTIF({qr("측정 회차")},">0")'
n_hit = f'COUNTIFS({qr("측정 회차")},">0",{qr("묘책 언급 회차")},">0")'
x_meas = f'COUNTIFS({qr("브랜드")},"X",{qr("측정 회차")},">0")'
x_hit = f'COUNTIFS({qr("브랜드")},"X",{qr("측정 회차")},">0",{qr("묘책 언급 회차")},">0")'
o_meas = f'COUNTIFS({qr("브랜드")},"O",{qr("측정 회차")},">0")'
o_hit = f'COUNTIFS({qr("브랜드")},"O",{qr("우리 사이트 인용 회차")},">0")'
o_runs = f'COUNTIF({rng("질문에 묘책 이름")},"O")'
o_runs_hit = f'COUNTIFS({rng("질문에 묘책 이름")},"O",{rng("우리 사이트 인용")},"Y")'
cards = [
    ("① 커버리지 (기준서 산식)", f"={n_hit}/{n_meas}", f'={n_hit}&" / "&{n_meas}&"개 질문"',
     "묘책이 한 번이라도 나온 질문 ÷ 측정한 질문. 나온 6개는 전부 질문에 「묘책」을 직접 넣은 질문이라 답이 이름을 따라 말한 거예요."),
    ("② 진짜 커버리지 (이름 없는 질문)", f"={x_hit}/{x_meas}", f'={x_hit}&" / "&{x_meas}&"개 질문"',
     "질문에 묘책 이름이 없는데 ChatGPT가 스스로 묘책을 꺼낸 비율. 매달 올려야 하는 진짜 목표 숫자예요."),
    ("③ 이름으로 물었을 때 우리 사이트 인용", f"={o_hit}/{o_meas}",
     f'={o_hit}&" / "&{o_meas}&"개 질문 · "&{o_runs_hit}&" / "&{o_runs}&"회"',
     "「묘책사주 어때?」처럼 이름을 넣어 물었을 때 myocheck.kr(칼럼 포함)이 답의 출처 링크로 붙은 비율."),
]
for label, val, sub, note in cards:
    line([label, val, sub, note], bold_cols=(0, 1), fmt={1: "0.0%"}, wrap_cols=(0, 3), height=44, fill=fill_card)
    d.cell(row=row - 1, column=2).font = Font(name=F, size=16, bold=True)
    d.cell(row=row - 1, column=2).alignment = Alignment(horizontal="center", vertical="center")

h2("측정 개요")
header(["항목", "값", "", "설명"])
over = [
    ("질문 수 (세트)", f"=COUNTA({QD}!$A$2:$A${NQ})", ""),
    ("측정한 질문 수", f"={n_meas}", ""),
    ("미측정 질문", f'=COUNTIF({qr("상태")},"미측정")', "S04·S05 — S04-1에서 ChatGPT 계정이 로그아웃돼 174회에서 끝냄"),
    ("측정 회차 합계", f"=COUNTA({rng('실행ID')})", ""),
    ("  └ gpt-6 회차", f'=COUNTIF({rng("모델")},"gpt-6")', "I01~I16·A01 (처음 51회)"),
    ("  └ gpt-6-mini 회차", f'=COUNTIF({rng("모델")},"gpt-6-mini")', "무료 계정이라 중간부터 가벼운 모델로 바뀜"),
    ("ChatGPT가 웹검색을 한 회차", f'=COUNTIF({rng("웹검색 함")},"Y")', "검색을 해야 우리 글이 출처로 붙을 수 있어요"),
    ("이름 없는 질문 중 빈칸 (= 다음 달 글감 후보)", f'=COUNTIFS({qr("브랜드")},"X",{qr("상태")},"빈칸")', "질문세트 시트 「다음 달 글감 순위」 칸에 표시"),
    ("  └ 1순위: ChatGPT가 검색까지 한 질문", f'=COUNTIF({qr("다음 달 글감 순위")},"1순위*")', "글이 생기면 출처로 잡힐 여지가 있는 질문"),
    ("  └ 2순위: 검색 없이 기억으로 답한 질문", f'=COUNTIF({qr("다음 달 글감 순위")},"2순위*")', "[추정] 바깥 출처가 쌓여야 움직이는 질문"),
]
for k, f_, note in over:
    line([k, f_, "", note], bold_cols=(1,), wrap_cols=(0, 3))

# 출처 도메인 상위 10
cnt, qset = Counter(), {}
for v in R.values():
    for dd in {dom(u) for u in v["cited"]} - {IMG}:
        if ours(dd):
            continue
        cnt[dd] += 1
        qset.setdefault(dd, set()).add(v["ID"])
top = [dd for dd, _ in sorted(cnt.items(), key=lambda x: (-x[1], -len(qset[x[0]]), x[0]))[:10]]
DOMNOTE = {
    "creatrip.com": "크리에이트립 — 외국인 대상 한국 여행·체험 소개 (사주카페 소개 글)",
    "google.com": "구글 지도 링크 (오프라인 사주카페 길찾기)",
    "semojum.com": "점집·사주 정보 모음 사이트",
    "blog.naver.com": "네이버 블로그 후기",
    "kmong.com": "크몽 — 사주 상담 판매",
    "cm.asiae.co.kr": "아시아경제 기사",
    "sajuai.io": "AI 사주 서비스 (사주AI 소연)",
    "witchnextdoor.co": "이웃집 마녀 — 온라인 사주·타로",
    "datepop.co.kr": "데이트팝 — 사주카페 예약·후기",
    "starfortune.kr": "사주·타로 정보",
    "modernwitch.net": "사주·타로 정보",
}
h2("출처 도메인 상위 10 (답에 링크로 붙은 곳 = 후기·글을 실을 후보)")
header(["도메인", "회차", "질문 수", "어떤 곳"])
for dd in top:
    line([dd,
          f'=COUNTIF({rng("계산용_인용도메인")},"*|{dd}|*")',
          f'=SUMPRODUCT(({qr("측정 회차")}>0)*(COUNTIFS({rng("질문ID")},{QD}!$A$2:$A${NQ},{rng("계산용_인용도메인")},"*|{dd}|*")>0))',
          DOMNOTE.get(dd, "")], wrap_cols=(3,))
para("이미지 주소(images.openai.com)와 우리 사이트(myocheck.kr)는 뺐어요. 우리 사이트는 위 ③에 따로 있어요. "
     "도메인 설명은 클로드가 답 원문과 주소를 보고 적은 거예요.", f_note)

# 경쟁사
h2("경쟁사가 답에 나온 횟수 (모델별로 나눠서 셈)")
header(["경쟁사", "gpt-6 회차", "gpt-6-mini 회차", "합계 회차"])
ccount = Counter()
for v in R.values():
    for c_, al in COMPS.items():
        if any(a in v["text"] for a in al):
            ccount[c_] += 1
clist = [c_ for c_, n in ccount.most_common() if n > 0]
for c_ in clist:
    r_ = row
    line([c_,
          f'=COUNTIFS({rng("계산용_경쟁사")},"*|{c_}|*",{rng("모델")},"gpt-6")',
          f'=COUNTIFS({rng("계산용_경쟁사")},"*|{c_}|*",{rng("모델")},"gpt-6-mini")',
          f"=B{r_}+C{r_}"], wrap_cols=())
para("모델이 질문 묶음과 겹쳐요: I01~I16·A01은 gpt-6, 나머지는 전부 gpt-6-mini라서 「모델 차이」와 「질문 종류 차이」를 이번 측정만으로는 나눌 수 없어요. "
     "청월당 7회 중 6회(A07·C10)는 질문에 「청월당」이 들어 있어 답이 따라 말한 거예요. 스스로 꺼낸 건 A15 1회예요. "
     "목록은 측정 전에 정한 곳들 + 답을 읽고 추가한 2곳(이웃집 마녀·사주AI 소연)이고, 1회 이상 나온 곳만 표에 보여요. 신점·타로는 업체가 아닌 분야라, 오프라인 사주카페는 온라인 경쟁사가 아니라서 뺐어요.", f_note, 58)

h2("이번 측정에서 보인 것 (원본데이터에서 직접 확인)")
finds = [
    "1. 이름 없는 질문 52개에서 묘책은 한 번도 안 나왔어요. 지금 ChatGPT는 묘책을 「추천할 곳」으로 모르는 상태예요.",
    "2. 「묘책사주 어때?」처럼 이름을 넣으면 ChatGPT가 검색해서 column.myocheck.kr을 출처로 붙였어요 (D01·D02·D03은 3회 모두). 칼럼 사이트는 이름으로는 찾힌다는 뜻이에요.",
    "3. D05 「묘책사주 2027 신년운세 어때?」는 3회 모두 검색을 안 했어요. 3회 중 2회는 「묘책」을 사람 이름처럼 읽어 \"묘책의 2027년 운세\"를 봐주겠다며 생년월일을 물었어요. 신년운세와 묘책을 이어 주는 글이 아직 ChatGPT에 안 잡힌다는 신호예요.",
    "4. 「~잘 보는 곳 추천해줘」 질문(A그룹)에 ChatGPT는 오프라인 사주카페(구글 지도 링크·데이트팝·크리에이트립)와 이웃집 마녀·크몽·포스텔러·사주AI 소연을 꺼냈어요. A그룹 57회 중 28회는 예산이나 지역을 되물었어요.",
    "5. 정보형(I그룹 16개)·상황형(S그룹 3개) 질문은 3회 모두 검색 없이 답했어요. [추정] 이런 질문은 칼럼 한 편만으로는 바로 인용되기 어렵고, 기준서 6번처럼 바깥 출처(후기·기사)가 쌓여야 움직일 가능성이 커요.",
]
for t_ in finds:
    para(t_, f_base, 44)

h2("산식 · 기간 · 세트 버전")
foot = [
    "① 커버리지 = 묘책이 1회 이상 나온 질문 수 ÷ 측정한 질문 수 (docs/geo-standard.md 4장 ① 산식 그대로)",
    "② 진짜 커버리지 = 같은 산식을 「질문에 묘책 이름이 없는 질문」에만 적용 (이름을 넣은 질문은 답이 이름을 따라 말하는 것이라 빼요)",
    "③ 우리 사이트 인용 = 이름을 넣은 질문 중 답의 출처 링크에 myocheck.kr(칼럼 포함)이 1회 이상 붙은 질문 수 ÷ 측정한 이름 질문 수",
    "묘책 언급 판정 = 답 본문에 「묘책사주·묘책 사주·myocheck·MYO:CHECK·재회막차」 중 하나가 있으면 언급 (「묘책」 한 단어만은 일반 단어라 세지 않음)",
    "웹검색 판정 = 그 회차 답에 검색 결과 목록이 하나라도 기록됐으면 「검색함」",
    f"기간 = {period} · 측정 도구 = ChatGPT 임시 채팅(메모리 미사용), 새 계정, 질문당 3회",
    "세트 버전 = 질문세트 v1.0 (60문) · 미측정 = S04, S05",
    "제외한 자료 = 첫 계정으로 잰 19회는 메모리 오염 가능성이 있어 넣지 않았어요 (I07-1 답에서 「묘책 콘텐츠 만드는 거야?」라고 되물음)",
    "다음 측정 = 2026-11 같은 세트·같은 방식으로 반복. 1~2회 결과로 판단하지 않고 2~3회 추세로 봐요.",
]
for t_ in foot:
    para(t_, f_note, 28)
d.sheet_view.showGridLines = False
d.freeze_panes = "A3"
for w_ in (d, ws_q):
    w_.sheet_properties.pageSetUpPr.fitToPage = True
    w_.page_setup.fitToWidth = 1
    w_.page_setup.fitToHeight = 0
    w_.page_setup.orientation = "landscape"

wb.save(OUT)
print("saved", OUT, "raw rows", NR - 1, "questions", NQ - 1, "top domains", top, "comps", clist)
